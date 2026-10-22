import shutil
import zipfile

import openpyxl
import pandas as pd
import pytest

from mortality.config import Paths
from mortality.ingest import cnsf_vi

HEADERS = {
    "Comisiones": [
        "EDAD",
        "PLAN DE LA POLIZA",
        "MODALIDAD DE LA POLIZA",
        "MONEDA",
        "ENTIDAD ",
        "SEXO",
        "FORMA DE VENTA",
        "TIPO DIVIDENDO",
        "NUMERO DE ASEGURADOS",
        "PRIMA CEDIDA",
        "COMISIONES DIRECTAS",
        "FONDO DE INVERSIÓN",
        "FONDO DE ADMINISTRACION",
        "MONTO DE DIVIDENDOS",
        "MONTO DE RESCATE",
    ],
    "Emision": [
        "EDAD",
        "COBERTURA",
        "PLAN DE LA POLIZA",
        "MODALIDAD DE LA POLIZA",
        "MONEDA",
        "ENTIDAD ",
        "SEXO",
        "FORMA DE VENTA",
        "NUMERO DE ASEGURADOS",
        "PRIMA EMITIDA",
        "SUMA ASEGURADA",
    ],
    "Siniestros": [
        "EDAD",
        "COBERTURA",
        "ENTIDAD",
        "CAUSA DEL SINIESTRO",
        "PLAN DE LA POLIZA",
        "MODALIDAD DE LA POLIZA",
        "SEXO",
        "NUMERO DE SINIESTROS",
        "MONTO RECLAMADO",
        "VENCIMIENTOS",
        "MONTO PAGADO",
        "MONTO DE REASEGURO",
    ],
}
ROWS = {
    "Comisiones": [
        (
            40,
            "Temporal",
            "Tradicional",
            "Nacional",
            "Jalisco",
            "Masculino",
            "Agentes",
            "Sin",
            10,
            0,
            5.5,
            0,
            0,
            0,
            0,
        ),
        (
            41,
            "Vitalicio",
            "Tradicional",
            "Nacional",
            "Jalisco",
            "Femenino",
            "Agentes",
            "Sin",
            4,
            0,
            2.0,
            0,
            0,
            0,
            0,
        ),
        (
            "No disponible",
            "No disponible",
            "No disponible",
            "Nacional",
            "Jalisco",
            "No disponible",
            "Agentes",
            "Sin",
            3,
            0,
            1.0,
            0,
            0,
            0,
            0,
        ),
    ],
    "Emision": [
        (
            40,
            "Fallecimiento",
            "Temporal",
            "Tradicional",
            "Nacional",
            "Jalisco",
            "Masculino",
            "Agentes",
            10,
            100.5,
            1_000_000,
        ),
        (
            41,
            "Fallecimiento",
            "Vitalicio",
            "Tradicional",
            "Nacional",
            "Jalisco",
            "Femenino",
            "Agentes",
            4,
            50,
            400_000,
        ),
    ],
    "Siniestros": [
        (
            40,
            "Fallecimiento",
            "Jalisco",
            "INFARTO",
            "Temporal",
            "Tradicional",
            "Masculino",
            2,
            200_000,
            0,
            200_000,
            0,
        ),
        (
            "No disponible",
            "Fallecimiento",
            "Jalisco",
            "SEPSIS",
            "No disponible",
            "No disponible",
            "No disponible",
            1,
            50_000,
            0,
            50_000,
            0,
        ),
        (
            41,
            "Gastos funerarios",
            "Jalisco",
            "OTRA",
            "Vitalicio",
            "Tradicional",
            "Femenino",
            1,
            10_000,
            0,
            10_000,
            0,
        ),
    ],
}


def write_workbook(path, header_override=None, skip_sheet=None):
    wb = openpyxl.Workbook()
    wb.active.title = "Indice"
    for name, header in HEADERS.items():
        if name == skip_sheet:
            continue
        ws = wb.create_sheet(name)
        ws.append(["Indice", "TITLE"])
        ws.append(header_override if header_override and name == "Siniestros" else header)
        for row in ROWS[name]:
            ws.append(list(row))
        ws.append([None] * len(header))  # blank trailing row
    wb.save(path)


@pytest.fixture
def paths(tmp_path):
    return Paths(root=tmp_path, data=tmp_path / "data").ensure()


def test_fold_removes_accents_and_spaces():
    assert cnsf_vi.fold("  Fondo de  Inversión ") == "FONDO DE INVERSION"
    assert cnsf_vi.fold(None) == ""


def test_find_files_matches_both_naming_styles(tmp_path):
    for name in (
        "2024 Vida Individual Bases.xlsx",
        "Vida_Ind_2018.zip",
        "Vida_ind_2019.zip",
        "Vida Individual 2020.zip",
        "notes.txt",
    ):
        (tmp_path / name).write_bytes(b"x")
    assert sorted(cnsf_vi.find_files(tmp_path)) == [2018, 2019, 2020, 2024]


def test_read_sheet_cleans_values_and_keeps_every_row(tmp_path):
    book = tmp_path / "w.xlsx"
    write_workbook(book)
    frame, control = cnsf_vi.read_sheet(book, cnsf_vi.SHEET_BY_NAME["Comisiones"], year=2024)
    assert len(frame) == 3  # the blank trailing row is skipped, nothing else is
    assert list(frame["sex"]) == ["M", "F", pd.NA]
    assert frame["age"].isna().tolist() == [False, False, True]
    assert frame["plan"].isna().tolist() == [False, False, True]
    assert (frame["year"] == 2024).all()
    total = control.set_index("column").loc["insured", "total"]
    assert total == 17 == frame["insured"].sum()
    assert control.set_index("column").loc["insured", "non_numeric"] == 0


def test_control_totals_are_independent_of_cleaning(tmp_path):
    book = tmp_path / "w.xlsx"
    write_workbook(book)
    frame, control = cnsf_vi.read_sheet(book, cnsf_vi.SHEET_BY_NAME["Siniestros"], year=2024)
    ctl = control.set_index("column")
    assert ctl.loc["claims", "total"] == 4
    assert ctl.loc["amount_claimed", "total"] == 260_000
    assert frame["claims"].sum() == ctl.loc["claims", "total"]


def test_unrecognised_sex_and_age_labels_are_counted_not_hidden(tmp_path):
    book = tmp_path / "w.xlsx"
    write_workbook(book)
    wb = openpyxl.load_workbook(book)
    ws = wb["Siniestros"]
    ws.append(
        [41.5, "Fallecimiento", "Jalisco", "X", "Temporal", "Tradicional", "Otro", 1, 1, 0, 1, 0]
    )
    ws.append(
        ["adulto", "Fallecimiento", "Jalisco", "X", "Temporal", "Tradicional", None, 1, 1, 0, 1, 0]
    )
    ws.append(  # a blank age next to 'No disponible' sex is simply missing, not unexpected
        [
            None,
            "Fallecimiento",
            "Jalisco",
            "X",
            "Temporal",
            "Tradicional",
            "No disponible",
            1,
            1,
            0,
            1,
            0,
        ]
    )
    wb.save(book)
    frame, control = cnsf_vi.read_sheet(book, cnsf_vi.SHEET_BY_NAME["Siniestros"], year=2024)
    ctl = control.set_index("column")["total"]
    assert ctl["_unexpected_sex"] == 2  # 'Otro' and a blank cell
    assert ctl["_unexpected_age"] == 2  # 41.5 and 'adulto'; a blank age is just missing
    assert frame["sex"].isna().sum() == 4  # 'Otro', blank, and two legitimate 'No disponible'
    assert frame["age"].isna().sum() == 4  # the 'No disponible' row, 41.5, 'adulto', blank


def test_changed_header_is_refused(tmp_path):
    book = tmp_path / "w.xlsx"
    bad = list(HEADERS["Siniestros"])
    bad[7] = "NUMERO DE EVENTOS"
    write_workbook(book, header_override=bad)
    with pytest.raises(cnsf_vi.CNSFVIError, match="unexpected header"):
        cnsf_vi.read_sheet(book, cnsf_vi.SHEET_BY_NAME["Siniestros"], year=2024)


def test_missing_sheet_is_refused(tmp_path):
    book = tmp_path / "w.xlsx"
    write_workbook(book, skip_sheet="Emision")
    with pytest.raises(cnsf_vi.CNSFVIError, match="no sheet"):
        cnsf_vi.read_sheet(book, cnsf_vi.SHEET_BY_NAME["Emision"], year=2024)


def test_import_folder_copies_workbooks_and_extracts_databases_safely(tmp_path, paths):
    source = tmp_path / "downloads"
    source.mkdir()
    write_workbook(source / "2024 Vida Individual Bases.xlsx")
    with zipfile.ZipFile(source / "Vida_ind_2019.zip", "w") as zf:
        zf.writestr("../../escape/Vida_Ind_2019.mdb", b"not really access")
    imported = cnsf_vi.import_folder(source, paths=paths, log=lambda _: None)
    assert imported[2024].parent == paths.raw / "cnsf_vi"
    assert imported[2019].name == "2019_vida_individual.mdb"
    assert imported[2019].parent == paths.raw / "cnsf_vi"  # nothing escaped the target folder
    assert not (tmp_path / "escape").exists()
    manifest = pd.read_csv(paths.manifest)
    assert len(manifest) == 2


def test_zip_with_two_databases_is_refused(tmp_path, paths):
    source = tmp_path / "downloads"
    source.mkdir()
    with zipfile.ZipFile(source / "Vida_Ind_2018.zip", "w") as zf:
        zf.writestr("a.mdb", b"1")
        zf.writestr("b.accdb", b"2")
    with pytest.raises(cnsf_vi.CNSFVIError, match="one Access database"):
        cnsf_vi.import_folder(source, paths=paths, log=lambda _: None)


def test_process_year_caches_and_combine_concatenates(tmp_path, paths):
    book = tmp_path / "w.xlsx"
    write_workbook(book)
    (paths.raw / "cnsf_vi").mkdir(parents=True, exist_ok=True)
    for year in (2023, 2024):
        shutil.copyfile(book, paths.raw / "cnsf_vi" / f"{year}_vida_individual.xlsx")
    messages = []
    written = cnsf_vi.process(paths=paths, legacy=False, log=messages.append)
    names = {p.name for p in written}
    assert {"exposure.parquet", "issued.parquet", "claims.parquet", "control.parquet"} <= names
    claims = pd.read_parquet(paths.processed / "cnsf_vi" / "claims.parquet")
    assert sorted(claims["year"].unique()) == [2023, 2024]
    assert len(claims) == 6
    cached = []
    cnsf_vi.process(paths=paths, legacy=False, log=cached.append)
    assert any("cached" in m for m in cached) and not any("reading" in m for m in cached)


def test_lookup_catalog_stacks_code_tables_whatever_their_headers():
    tables = {
        1: pd.DataFrame({"mod_pol": ["1", "2", ""], "Descripción": ["Temporal", "Vitalicio", ""]}),
        5: pd.DataFrame({"FORM_VENTA": ["01"], "DESCRIPCION": ["Agente (persona física)"]}),
    }
    catalog = cnsf_vi.lookup_catalog(tables)
    assert catalog.to_dict("records") == [
        {"catalog": 1, "code": "1", "label": "Temporal"},
        {"catalog": 1, "code": "2", "label": "Vitalicio"},
        {"catalog": 5, "code": "01", "label": "Agente (persona física)"},
    ]


def test_clean_legacy_decodes_codes_even_when_some_are_missing():
    catalog = pd.DataFrame(
        {
            "catalog": [1, 1, 4, 4, 5],
            "code": ["1", "2", "1", "2", "01"],
            "label": ["Temporal", "Vitalicio", "Aguascalientes", "Baja California", "Agente"],
        }
    )
    issued = pd.DataFrame(
        {
            "MOD_POL": [1, 2],
            "ENTIDAD": [1.0, None],  # floats because one code is missing
            "FORM_VENTA": [1, 1],
            "EDAD": [40, 41],
            "SEXO": ["M", "F"],
            "ASEGURADOS": [5, 6],
        }
    )
    claims = pd.DataFrame(
        {
            "MOD_POL": [1],
            "ENTIDAD": [None],
            "FORM_VENTA": [None],
            "EDAD": [40.0],
            "SEXO": ["m"],
            "SINIESTROS": [2],
        }
    )
    # 2018-style claims: no ENTIDAD column, only the state where the claim occurred
    claims = claims.drop(columns="ENTIDAD").assign(ENT_OCURSIN=[2])
    i, c = cnsf_vi.clean_legacy(issued, claims, catalog, year=2019)
    assert list(c["state"]) == ["Baja California"]
    claims = claims.assign(ENTIDAD=[None], ENT_OCURSIN=[1])
    _, c = cnsf_vi.clean_legacy(issued, claims, catalog, year=2019)
    assert list(i["plan"]) == ["Temporal", "Vitalicio"]
    assert i["state"].tolist() == ["Aguascalientes", pd.NA]
    assert list(i["channel"]) == ["Agente", "Agente"]
    assert list(c["sex"]) == ["M"]
    assert list(c["state"]) == ["Aguascalientes"]  # empty ENTIDAD filled from the occurrence state
    assert (i["year"] == 2019).all() and i["age"].dtype == "Int16"
