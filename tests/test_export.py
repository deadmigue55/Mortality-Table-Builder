"""The Excel workbook: live formulas, checked against the Python commutation functions."""

from __future__ import annotations

import sqlite3
import warnings
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from mortality import export, lifetable

BUNDLE = Path(__file__).resolve().parent.parent / "app" / "data" / "demo.sqlite"
pytestmark = pytest.mark.skipif(not BUNDLE.exists(), reason="demo bundle not built")

KEYS = {"Mexico 2023, men": "M", "Mexico 2023, women": "F", "CNSF 2013 (regulatory)": "cnsf"}


@pytest.fixture(scope="module")
def con():
    c = sqlite3.connect(BUNDLE)
    yield c
    c.close()


def calculate(path: Path) -> dict[str, float]:
    formulas = pytest.importorskip("formulas")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, value in solution.items():
        if "CALCULATOR'!" in key.upper() and ":" not in key.split("!")[-1]:
            cell = key.split("!")[-1].upper()
            out[cell] = value.value[0, 0]
    return out


def expected(con, table: str, age: int, interest: float, term: int) -> dict[str, float]:
    qx = export._qx_curve(con, KEYS[table])
    life = lifetable.build_life_table(np.arange(len(qx)), qx)
    com = lifetable.Commutation.from_table(life, interest)
    lx = life["lx"].to_numpy()
    return {
        "B13": float(lx[age + 1 :].sum() / lx[age]),
        "B14": float(com.whole_life_insurance(age)),
        "B15": float(com.annuity_due(age)),
        "B16": float(com.term_insurance(age, term)),
        "B17": float(com.temporary_annuity_due(age, term)),
        "B18": float(com.term_premium(age, term)),
    }


@pytest.mark.parametrize(
    "inputs",
    [
        {"table": "Mexico 2023, men", "age": 40, "interest": 0.04, "term": 20},
        {"table": "CNSF 2013 (regulatory)", "age": 55, "interest": 0.05, "term": 10},
        {"table": "Mexico 2023, women", "age": 30, "interest": 0.03, "term": 25},
    ],
)
def test_the_spreadsheet_agrees_with_the_python_commutation_functions(con, tmp_path, inputs):
    path = tmp_path / "wb.xlsx"
    export.build_workbook(con, inputs).save(path)
    got = calculate(path)
    want = expected(con, inputs["table"], inputs["age"], inputs["interest"], inputs["term"])
    for cell, value in want.items():
        assert got[cell] == pytest.approx(value, rel=1e-9), cell


def test_the_tables_hold_formulas_not_pasted_values(con, tmp_path):
    path = tmp_path / "wb.xlsx"
    export.write_workbook(con, path)
    wb = load_workbook(path)
    ws = wb["Table_M"]
    assert ws["C2"].value == export.RADIX
    for cell in ("C3", "D3", "E3", "F3", "G3", "H3"):
        assert str(ws[cell].value).startswith("="), cell
    assert "Calculator" in str(ws["E3"].value)  # driven by the interest-rate input


def test_the_workbook_has_every_sheet(con, tmp_path):
    path = tmp_path / "wb.xlsx"
    export.write_workbook(con, path)
    names = load_workbook(path).sheetnames
    for sheet in ("README", "Calculator", "Table_M", "Table_F", "Table_CNSF", "Rate_Card",
                  "Excess_Mortality", "Annuity_Stress", "Table_Comparison"):  # fmt: skip
        assert sheet in names


def test_the_workbook_bytes_round_trip(con):
    data = export.workbook_bytes(con)
    assert data[:2] == b"PK"  # a zip container
    assert len(data) > 20_000
