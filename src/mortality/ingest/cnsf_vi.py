"""CNSF individual life insurance statistics (Sistema Estadístico, Vida Individual).

Every year the CNSF publishes, for the whole Mexican individual life market, the number of
insured lives and the death claims, by age, sex, plan, modality and state. This is the closest
public thing to an *industry experience study*: it describes the people who actually buy life
insurance, not the whole population.

Two formats exist and they are read separately, never mixed silently:

* **2021 to 2025, Excel workbooks** (``<year> Vida Individual Bases.xlsx``) with three sheets:

  - ``Comisiones``: policies exposed during the year (the **exposure**: insured lives by age,
    sex, plan, modality, state), with premium and commission amounts;
  - ``Emision``: premium and **sum assured** by coverage;
  - ``Siniestros``: **claims** by coverage, with the cause of the claim as text.

* **2018 to 2020, Microsoft Access databases** with coded columns, no cause of death and a
  single insured count in the issue table. Reading them needs the Windows Access ODBC driver,
  so a small PowerShell script (``access_export.ps1``) exports the tables to CSV. The legacy
  tables are kept apart (``legacy_*``) because their exposure definition differs.

Raw files are copied to ``data/raw/cnsf_vi/`` and logged in the SHA-256 manifest. Parsing the
workbooks is slow (minutes per file), so each year is cached as Parquet.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from mortality.config import Paths, get_paths
from mortality.ingest.download import register_local

SOURCE = "cnsf_vi"
PAGE_URL = "https://www.gob.mx/cnsf/acciones-y-programas/estadistica-de-seguros"
MODERN_YEARS = (2021, 2022, 2023, 2024, 2025)
LEGACY_YEARS = (2018, 2019, 2020)
UNKNOWN = "no disponible"
EXPORT_SCRIPT = Path(__file__).with_name("access_export.ps1")

_XLSX_NAME = re.compile(r"^(\d{4}) Vida Individual Bases\.xlsx$", re.IGNORECASE)
_ZIP_NAME = re.compile(r"^vida[_ ]ind(?:ividual)?[_ ](\d{4})\.zip$", re.IGNORECASE)
_DB_SUFFIXES = (".mdb", ".accdb")


class CNSFVIError(ValueError):
    """A workbook or database does not look like we expect (layout changed, or a bad file)."""


@dataclass(frozen=True)
class Sheet:
    """One worksheet: the header we expect (accent-folded), the names we give the columns."""

    sheet: str
    table: str
    header: tuple[str, ...]
    columns: tuple[str, ...]
    text: tuple[str, ...]
    counts: tuple[str, ...]
    amounts: tuple[str, ...]


SHEETS: tuple[Sheet, ...] = (
    Sheet(
        "Comisiones",
        "exposure",
        (
            "EDAD",
            "PLAN DE LA POLIZA",
            "MODALIDAD DE LA POLIZA",
            "MONEDA",
            "ENTIDAD",
            "SEXO",
            "FORMA DE VENTA",
            "TIPO DIVIDENDO",
            "NUMERO DE ASEGURADOS",
            "PRIMA CEDIDA",
            "COMISIONES DIRECTAS",
            "FONDO DE INVERSION",
            "FONDO DE ADMINISTRACION",
            "MONTO DE DIVIDENDOS",
            "MONTO DE RESCATE",
        ),
        (
            "age",
            "plan",
            "modality",
            "currency",
            "state",
            "sex",
            "channel",
            "dividend_type",
            "insured",
            "ceded_premium",
            "direct_commission",
            "investment_fund",
            "admin_fund",
            "dividends",
            "surrender_value",
        ),
        ("plan", "modality", "currency", "state", "sex", "channel", "dividend_type"),
        ("insured",),
        (
            "ceded_premium",
            "direct_commission",
            "investment_fund",
            "admin_fund",
            "dividends",
            "surrender_value",
        ),
    ),
    Sheet(
        "Emision",
        "issued",
        (
            "EDAD",
            "COBERTURA",
            "PLAN DE LA POLIZA",
            "MODALIDAD DE LA POLIZA",
            "MONEDA",
            "ENTIDAD",
            "SEXO",
            "FORMA DE VENTA",
            "NUMERO DE ASEGURADOS",
            "PRIMA EMITIDA",
            "SUMA ASEGURADA",
        ),
        (
            "age",
            "coverage",
            "plan",
            "modality",
            "currency",
            "state",
            "sex",
            "channel",
            "insured",
            "premium",
            "sum_assured",
        ),
        ("coverage", "plan", "modality", "currency", "state", "sex", "channel"),
        ("insured",),
        ("premium", "sum_assured"),
    ),
    Sheet(
        "Siniestros",
        "claims",
        (
            "EDAD",
            "COBERTURA",
            "ENTIDAD",
            "CAUSA DEL SINIESTRO",
            "PLAN DE LA POLIZA",
            "MODALIDAD DE LA POLIZA",
            "SEXO",
            "NUMERO DE SINIESTROS",
            "MONTO RECLAMADO",
            "VENCIMIENTOS",
            "MONTO PAGADO",
            "MONTO DE REASEGURO",
        ),
        (
            "age",
            "coverage",
            "state",
            "cause",
            "plan",
            "modality",
            "sex",
            "claims",
            "amount_claimed",
            "maturities",
            "amount_paid",
            "reinsured",
        ),
        ("coverage", "state", "cause", "plan", "modality", "sex"),
        ("claims",),
        ("amount_claimed", "maturities", "amount_paid", "reinsured"),
    ),
)
SHEET_BY_NAME = {s.sheet: s for s in SHEETS}

SEX_CODES = {"masculino": "M", "femenino": "F", "m": "M", "f": "F"}

# Legacy (Access) layout. Numeric columns are summed for the control totals.
LEGACY_TABLES = ("EMISION", "SINIESTROS")
# The code lists are small tables of their own in every year (code, description). Their numbers
# follow the 'Catálogo 1 .. 6' of the data dictionary.
LEGACY_LOOKUPS = {
    1: "MOD_POL",  # plan
    2: "STATUS_POL",  # policy status
    3: "STATUS_SIN",  # claim status
    4: "ENTIDAD",  # state
    5: "FORM_VENTA",  # channel
    6: "SUBT_SEG",  # product subtype
}


def fold(text: object) -> str:
    """Upper-case, accent-free, single-spaced text, for comparing headers."""
    decomposed = unicodedata.normalize("NFKD", "" if text is None else str(text))
    plain = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", plain).strip().upper()


# ── Importing the raw files ─────────────────────────────────────────────────


def find_files(folder: Path) -> dict[int, Path]:
    """The yearly files in a folder: workbooks for 2021+, zipped databases for 2018 to 2020."""
    found: dict[int, Path] = {}
    for path in sorted(folder.iterdir()):
        match = _XLSX_NAME.match(path.name) or _ZIP_NAME.match(path.name)
        if match and path.is_file():
            found[int(match.group(1))] = path
    return found


def import_folder(
    folder: Path, *, paths: Paths | None = None, log: Callable[[str], None] = print
) -> dict[int, Path]:
    """Copy the yearly files into ``data/raw/cnsf_vi`` and log them in the manifest.

    Workbooks are copied. A zip must hold exactly one Access database, which is extracted by
    file name only (nothing in the archive can write outside the target folder).
    """
    paths = paths or get_paths()
    paths.ensure()
    target = paths.raw / SOURCE
    target.mkdir(parents=True, exist_ok=True)
    imported: dict[int, Path] = {}
    for year, source in find_files(folder).items():
        if source.suffix.lower() == ".xlsx":
            dest = target / f"{year}_vida_individual.xlsx"
            if not dest.exists() or dest.stat().st_size != source.stat().st_size:
                log(f"  copying {source.name}")
                shutil.copyfile(source, dest)
        else:
            dest = _extract_database(source, year, target, log)
        register_local(dest, source=SOURCE, url=f"{PAGE_URL}#{year}", paths=paths)
        imported[year] = dest
    return imported


def _extract_database(archive: Path, year: int, target: Path, log: Callable[[str], None]) -> Path:
    with zipfile.ZipFile(archive) as zf:
        members = [m for m in zf.infolist() if m.filename.lower().endswith(_DB_SUFFIXES)]
        if len(members) != 1:
            raise CNSFVIError(
                f"{archive.name} should hold one Access database, found {len(members)}"
            )
        member = members[0]
        suffix = Path(member.filename).suffix.lower()
        dest = target / f"{year}_vida_individual{suffix}"
        if not dest.exists() or dest.stat().st_size != member.file_size:
            log(f"  extracting {member.filename} from {archive.name}")
            with zf.open(member) as src, dest.open("wb") as out:
                shutil.copyfileobj(src, out)
    return dest


# ── Modern workbooks (2021+) ────────────────────────────────────────────────


def _to_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    return None


def _unexpected_sex(value: object) -> bool:
    """True for a sex cell that is neither a known label nor 'No disponible'."""
    text = value.strip().lower() if isinstance(value, str) else None
    return text is None or (text not in SEX_CODES and text != UNKNOWN)


def _unexpected_age(value: object) -> bool:
    """True for an age cell that is not a whole non-negative number, 'No disponible' or blank.

    A blank age is a legitimate way to say 'missing' (2025 leaves it empty where earlier
    years write 'No disponible').
    """
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip().lower() != UNKNOWN
    number = _to_number(value)
    return number is None or number < 0 or number != round(number)


def read_sheet(path: Path, sheet: Sheet, *, year: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Stream one worksheet. Returns the cleaned frame and its control totals.

    The control totals are summed from the raw cells while streaming, independently of the
    pandas cleaning, so the warehouse checks can prove nothing was lost or changed. Two extra
    rows (``_unexpected_sex``, ``_unexpected_age``) count raw labels the cleaning would
    otherwise turn into missing values without a trace.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet.sheet not in wb.sheetnames:
            raise CNSFVIError(f"{path.name} has no sheet {sheet.sheet!r} (found {wb.sheetnames})")
        ws = wb[sheet.sheet]
        rows = ws.iter_rows(min_row=2, values_only=True)
        header = next(rows)
        got = tuple(fold(c) for c in header[: len(sheet.header)])
        if got != sheet.header:
            raise CNSFVIError(
                f"{path.name}/{sheet.sheet}: unexpected header.\n"
                f"  expected {sheet.header}\n  found    {got}"
            )
        width = len(sheet.columns)
        numeric = {
            name: i for i, name in enumerate(sheet.columns) if name in sheet.counts + sheet.amounts
        }
        totals = dict.fromkeys(numeric, 0.0)
        non_numeric = dict.fromkeys(numeric, 0)
        sex_at, age_at = sheet.columns.index("sex"), sheet.columns.index("age")
        bad_sex = bad_age = 0
        data: list[tuple] = []
        for row in rows:
            row = tuple(row[:width])
            if all(c is None for c in row):
                continue
            if len(row) < width:
                row = row + (None,) * (width - len(row))
            bad_sex += _unexpected_sex(row[sex_at])
            bad_age += _unexpected_age(row[age_at])
            for name, i in numeric.items():
                number = _to_number(row[i])
                if number is None:
                    non_numeric[name] += 1
                else:
                    totals[name] += number
            data.append(row)
    finally:
        wb.close()

    frame = clean_frame(pd.DataFrame(data, columns=list(sheet.columns)), sheet)
    frame.insert(0, "year", year)
    control = pd.DataFrame(
        {
            "format": "xlsx",
            "year": year,
            "sheet": sheet.sheet,
            "column": [*numeric, "_unexpected_sex", "_unexpected_age"],
            "rows": len(frame),
            "total": [*(totals[c] for c in numeric), float(bad_sex), float(bad_age)],
            "non_numeric": [*(non_numeric[c] for c in numeric), 0, 0],
        }
    )
    return frame, control


def clean_frame(frame: pd.DataFrame, sheet: Sheet) -> pd.DataFrame:
    """Typed columns: 'No disponible' becomes missing, sex becomes M/F, age a nullable integer."""
    out = frame.copy()
    for col in sheet.text:
        text = out[col].astype("string").str.strip()
        out[col] = text.mask(text.str.lower() == UNKNOWN)
    out["sex"] = out["sex"].str.lower().map(SEX_CODES).astype("string")
    age = pd.to_numeric(out["age"], errors="coerce")
    out["age"] = age.where(age == age.round()).astype("Int16")
    for col in sheet.counts + sheet.amounts:
        out[col] = pd.to_numeric(out[col], errors="coerce").astype("float64")
    return out


def process_year(
    year: int, path: Path, *, paths: Paths | None = None, force: bool = False, log=print
) -> None:
    """Parse the three sheets of one workbook into per-year Parquet files."""
    paths = paths or get_paths()
    cache = paths.processed / SOURCE / "years"
    cache.mkdir(parents=True, exist_ok=True)
    for sheet in SHEETS:
        out = cache / f"{sheet.table}_{year}.parquet"
        ctl = cache / f"control_{sheet.table}_{year}.parquet"
        if out.exists() and ctl.exists() and not force:
            log(f"  {year} {sheet.sheet}: cached")
            continue
        log(f"  {year} {sheet.sheet}: reading (slow)...")
        frame, control = read_sheet(path, sheet, year=year)
        frame.to_parquet(out, index=False)
        control.to_parquet(ctl, index=False)
        log(f"  {year} {sheet.sheet}: {len(frame):,} rows")


# ── Legacy Access databases (2018 to 2020) ──────────────────────────────────


def access_available() -> bool:
    """True on Windows, where the Access ODBC driver and PowerShell exist."""
    return sys.platform == "win32" and shutil.which("powershell") is not None


def export_table(database: Path, table: str, out_csv: Path) -> None:
    """Export one Access table to UTF-8 CSV through the Windows ODBC driver.

    The export goes to a temporary file that is renamed only when it is complete, so an
    interrupted run never leaves a half-written CSV behind.
    """
    partial = out_csv.with_name(out_csv.name + ".part")
    partial.unlink(missing_ok=True)
    command = [
        "powershell",
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(EXPORT_SCRIPT),
        "-Path",
        str(database),
        "-Table",
        table,
        "-Out",
        str(partial),
    ]
    result = subprocess.run(command, capture_output=True, text=True, check=False)
    if result.returncode != 0 or not partial.exists():
        partial.unlink(missing_ok=True)
        raise CNSFVIError(f"Could not export {table} from {database.name}: {result.stderr.strip()}")
    partial.replace(out_csv)


def lookup_catalog(tables: dict[int, pd.DataFrame]) -> pd.DataFrame:
    """Stack the code-list tables into one (catalog, code, label) frame.

    Each table has the code in its first column and the description in its second, whatever
    the headers say (they carry accents and differ in case between years).
    """
    out = []
    for number, frame in tables.items():
        for code, label in zip(frame.iloc[:, 0], frame.iloc[:, 1], strict=True):
            if pd.notna(code) and pd.notna(label) and str(code).strip() and str(label).strip():
                out.append(
                    {"catalog": number, "code": str(code).strip(), "label": str(label).strip()}
                )
    return pd.DataFrame(out, columns=["catalog", "code", "label"])


def clean_legacy(
    issued: pd.DataFrame, claims: pd.DataFrame, catalog: pd.DataFrame, *, year: int
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Decode the legacy codes (plan, state, sex) so they line up with the modern tables."""
    labels = {
        cat: dict(zip(g["code"].astype(str).str.lstrip("0"), g["label"], strict=True))
        for cat, g in catalog.groupby("catalog")
    }

    def decode(series: pd.Series, catalog_no: int) -> pd.Series:
        # Codes arrive as floats when some are missing (1.0), so go through integers.
        keys = pd.to_numeric(series, errors="coerce").round().astype("Int64").astype("string")
        return keys.map(labels.get(catalog_no, {})).astype("string")

    out = []
    for frame in (issued, claims):
        f = frame.copy()
        f.columns = [c.lower() for c in f.columns]
        f["plan"] = decode(f["mod_pol"], 1)
        # Claims carry the state of the policyholder in some years and only the state where the
        # claim occurred in others (2018), so take the first one that is filled in.
        state = pd.Series(pd.NA, index=f.index, dtype="string")
        for column in ("entidad", "ent_ocursin"):
            if column in f.columns:
                state = state.fillna(decode(f[column], 4))
        f["state"] = state
        f["channel"] = decode(f["form_venta"], 5)
        f["sex"] = (
            f["sexo"].astype("string").str.strip().str.upper().where(lambda s: s.isin(["M", "F"]))
        )
        age = pd.to_numeric(f["edad"], errors="coerce")
        f["age"] = age.where(age == age.round()).astype("Int16")
        f.insert(0, "year", year)
        out.append(f)
    return out[0], out[1]


def process_legacy(
    year: int, database: Path, *, paths: Paths | None = None, force: bool = False, log=print
) -> None:
    """Export a legacy database to CSV and write the decoded tables as Parquet."""
    paths = paths or get_paths()
    cache = paths.processed / SOURCE / "years"
    cache.mkdir(parents=True, exist_ok=True)
    final = [cache / f"legacy_{name}_{year}.parquet" for name in ("issued", "claims", "control")]
    if all(p.exists() for p in final) and not force:
        log(f"  {year} legacy: cached")
        return
    if not access_available():
        raise CNSFVIError("Reading the 2018-2020 Access files needs Windows and PowerShell.")
    csv_dir = paths.raw / SOURCE / "legacy_csv" / str(year)
    csv_dir.mkdir(parents=True, exist_ok=True)
    for table in (*LEGACY_TABLES, *LEGACY_LOOKUPS.values()):
        target = csv_dir / f"{table}.csv"
        if target.exists() and not force:
            continue
        log(f"  {year} legacy {table}: exporting...")
        export_table(database, table, target)
    issued = pd.read_csv(csv_dir / "EMISION.csv", low_memory=False)
    claims = pd.read_csv(csv_dir / "SINIESTROS.csv", low_memory=False)
    catalog = lookup_catalog(
        {
            number: pd.read_csv(csv_dir / f"{table}.csv", dtype=str, keep_default_na=False)
            for number, table in LEGACY_LOOKUPS.items()
        }
    )
    controls = []
    for name, frame in (("EMISION", issued), ("SINIESTROS", claims)):
        numeric = frame.select_dtypes("number")
        controls.append(
            pd.DataFrame(
                {
                    "format": "access",
                    "year": year,
                    "sheet": name,
                    "column": numeric.columns,
                    "rows": len(frame),
                    "total": numeric.sum().to_numpy(dtype=float),
                    "non_numeric": 0,
                }
            )
        )
    issued, claims = clean_legacy(issued, claims, catalog, year=year)
    issued.to_parquet(final[0], index=False)
    claims.to_parquet(final[1], index=False)
    pd.concat(controls, ignore_index=True).to_parquet(final[2], index=False)
    catalog.assign(year=year).to_parquet(cache / f"legacy_catalog_{year}.parquet", index=False)
    log(f"  {year} legacy: {len(issued):,} issue rows, {len(claims):,} claim rows")


# ── Putting it together ─────────────────────────────────────────────────────


def process(
    *, paths: Paths | None = None, force: bool = False, legacy: bool = True, log=print
) -> list[Path]:
    """Parse every imported file, then combine the years into the final Parquet tables."""
    paths = paths or get_paths()
    raw = paths.raw / SOURCE
    for year in MODERN_YEARS:
        workbook = raw / f"{year}_vida_individual.xlsx"
        if workbook.exists():
            process_year(year, workbook, paths=paths, force=force, log=log)
    if legacy:
        for year in LEGACY_YEARS:
            database = next(
                (
                    raw / f"{year}_vida_individual{s}"
                    for s in _DB_SUFFIXES
                    if (raw / f"{year}_vida_individual{s}").exists()
                ),
                None,
            )
            if database is not None:
                process_legacy(year, database, paths=paths, force=force, log=log)
    return combine(paths)


def combine(paths: Paths | None = None) -> list[Path]:
    """Concatenate the per-year Parquet files into one file per table."""
    paths = paths or get_paths()
    cache = paths.processed / SOURCE / "years"
    out_dir = paths.processed / SOURCE
    groups = {
        "exposure": "exposure_*.parquet",
        "issued": "issued_*.parquet",
        "claims": "claims_*.parquet",
        "legacy_issued": "legacy_issued_*.parquet",
        "legacy_claims": "legacy_claims_*.parquet",
        "legacy_catalog": "legacy_catalog_*.parquet",
    }
    written = []
    for name, pattern in groups.items():
        files = sorted(cache.glob(pattern))
        if not files:
            continue
        frame = pd.concat([pd.read_parquet(f) for f in files], ignore_index=True)
        path = out_dir / f"{name}.parquet"
        frame.to_parquet(path, index=False)
        written.append(path)
    controls = [pd.read_parquet(f) for f in sorted(cache.glob("control_*.parquet"))]
    controls += [pd.read_parquet(f) for f in sorted(cache.glob("legacy_control_*.parquet"))]
    if controls:
        path = out_dir / "control.parquet"
        pd.concat(controls, ignore_index=True).to_parquet(path, index=False)
        written.append(path)
    return written
