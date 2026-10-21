"""An Excel workbook for actuaries who live in spreadsheets.

The life tables are written with **live formulas**, not pasted values: ``lx``, ``dx`` and the
commutation columns ``Dx, Nx, Cx, Mx`` are Excel formulas driven by one interest-rate cell, and a
calculator sheet prices insurance and annuities from them. Anyone can click a cell and audit the
arithmetic, or change the interest rate and watch every premium move. A test recomputes the
workbook with an independent formula engine and compares it with the Python commutation
functions, so the spreadsheet is checked against the code, not trusted.

Only the tables the dashboard bundle holds are used, so the workbook can be built from the demo
bundle as well as from the full warehouse.
"""

from __future__ import annotations

import io
import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FIRST_ROW = 2
LAST_AGE = 110
LAST_ROW = FIRST_ROW + LAST_AGE  # row of age 110
RADIX = 100_000
TABLES = (
    ("Table_M", "Mexico 2023, men"),
    ("Table_F", "Mexico 2023, women"),
    ("Table_CNSF", "CNSF 2013 (regulatory)"),
)
HEADER_FILL = PatternFill("solid", fgColor="0B0B0B")
INPUT_FILL = PatternFill("solid", fgColor="E8F1FC")
OUTPUT_FILL = PatternFill("solid", fgColor="F1F0EC")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
INPUT_FONT = Font(color="1F4FBF", bold=True)  # the convention: blue is an input


def _header(ws, columns: list[str], row: int = 1) -> None:
    for i, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font = WHITE_BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(name) + 3)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _qx_curve(con: sqlite3.Connection, which: str) -> np.ndarray:
    if which == "cnsf":
        sql = (
            "SELECT age, qx FROM cnsf_rates WHERE \"table\" = 'CNSF M 2013' AND variant = 'base' "
            "ORDER BY age"
        )
    else:
        sql = f"SELECT age, qx FROM graduated_qx WHERE year = 2023 AND sex = '{which}' ORDER BY age"
    q = pd.read_sql(sql, con).set_index("age")["qx"].reindex(range(LAST_AGE + 1)).ffill()
    q.iloc[-1] = 1.0
    return q.to_numpy()


def _table_sheet(wb: Workbook, name: str, label: str, qx: np.ndarray) -> None:
    ws = wb.create_sheet(name)
    _header(ws, ["Age x", "qx", "lx", "dx", "Dx", "Nx", "Cx", "Mx"])
    ws["J1"], ws["J2"] = label, "Edit the interest rate on the Calculator sheet."
    ws["J1"].font = Font(bold=True)
    for age in range(LAST_AGE + 1):
        r = FIRST_ROW + age
        ws.cell(row=r, column=1, value=age)
        ws.cell(row=r, column=2, value=float(qx[age])).number_format = "0.000000"
        ws.cell(row=r, column=3, value=RADIX if age == 0 else f"=C{r - 1}*(1-B{r - 1})")
        ws.cell(row=r, column=4, value=f"=C{r}*B{r}")
        ws.cell(row=r, column=5, value=f"=C{r}*(1/(1+Calculator!$B$5))^A{r}")
        ws.cell(row=r, column=6, value=f"=SUM(E{r}:E${LAST_ROW})")
        ws.cell(row=r, column=7, value=f"=D{r}*(1/(1+Calculator!$B$5))^(A{r}+1)")
        ws.cell(row=r, column=8, value=f"=SUM(G{r}:G${LAST_ROW})")
        for col in (3, 4, 5, 6, 7, 8):
            ws.cell(row=r, column=col).number_format = "#,##0.00"
    ws.column_dimensions["J"].width = 46


def _rng(sheet: str, column: str) -> str:
    return f"{sheet}!${column}${FIRST_ROW}:${column}${LAST_ROW}"


def _pick(column: str, position: str) -> str:
    """The value of a column of the chosen table at a position (1 = age 0): one INDEX per table,
    chosen by the table number in B9."""
    parts = ", ".join(f"INDEX({_rng(sheet, column)}, {position})" for sheet, _ in TABLES)
    return f"CHOOSE($B$9, {parts})"


def _calculator(wb: Workbook, defaults: dict | None = None) -> None:
    ws = wb.create_sheet("Calculator", 1)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["D"].width = 28
    ws["A1"], ws["A1"].font = "Actuarial calculator", Font(bold=True, size=14)
    ws["A2"] = "Blue cells are inputs. Everything else is a live formula on the Table_ sheets."
    d = {"table": TABLES[0][1], "age": 40, "interest": 0.04, "term": 20, "sum_assured": 1_000_000}
    d.update(defaults or {})
    inputs = [
        ("Mortality table", d["table"]),
        ("Issue age x", d["age"]),
        ("Interest rate i", d["interest"]),
        ("Term n (years)", d["term"]),
        ("Sum assured (MXN)", d["sum_assured"]),
    ]
    for offset, (label, value) in enumerate(inputs):
        r = 3 + offset
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=value)
        cell.font, cell.fill = INPUT_FONT, INPUT_FILL
    ws["B5"].number_format = "0.00%"
    ws["B7"].number_format = "#,##0"
    for offset, (_, label) in enumerate(TABLES):
        ws.cell(row=3 + offset, column=4, value=label)  # the list behind the drop-down
    ws["D2"] = "Tables"
    ws["A9"], ws["B9"] = "Table number (helper)", "=MATCH(B3, D3:D5, 0)"
    ws["A10"], ws["B10"] = "Position of age x (helper)", "=B4+1"
    ws["A11"], ws["B11"] = "Position of age x + n (helper)", f"=MIN(B4+B6,{LAST_AGE})+1"
    check = DataValidation(type="list", formula1="=$D$3:$D$5", allow_blank=False)
    ages = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
    ws.add_data_validation(check)
    ws.add_data_validation(ages)
    check.add("B3")
    ages.add("B4")
    dx, nx, mx = _pick("E", "$B$10"), _pick("F", "$B$10"), _pick("H", "$B$10")
    nxn, mxn = _pick("F", "$B$11"), _pick("H", "$B$11")
    lx = _pick("C", "$B$10")
    ex_sum = ", ".join(
        f'SUMIF({_rng(sheet, "A")}, ">"&$B$4, {_rng(sheet, "C")})' for sheet, _ in TABLES
    )
    out = [
        ("Curtate life expectancy e(x), years", f"=CHOOSE($B$9, {ex_sum})/({lx})"),
        ("Whole-life insurance A(x), per unit", f"=({mx})/({dx})"),
        ("Life annuity-due a(x), per unit a year", f"=({nx})/({dx})"),
        ("Term insurance A(x:n), per unit", f"=(({mx})-({mxn}))/({dx})"),
        ("Temporary annuity-due a(x:n)", f"=(({nx})-({nxn}))/({dx})"),
    ]
    start = 13
    ws.cell(row=start - 1, column=1, value="Results").font = Font(bold=True, size=12)
    for offset, (label, formula) in enumerate(out):
        r = start + offset
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=formula)
        c.fill, c.number_format = OUTPUT_FILL, "0.0000"
    ws["B13"].number_format = "0.0"
    r = start + len(out)
    ws.cell(row=r, column=1, value="Net annual premium rate, per unit of sum assured")
    ws.cell(row=r, column=2, value="=B16/B17").number_format = "0.000000"
    ws.cell(row=r + 1, column=1, value="Net annual premium, MXN")
    ws.cell(row=r + 1, column=2, value=f"=B{r}*B7").number_format = "#,##0"
    ws.cell(row=r + 2, column=1, value="Net premium per 1,000 of sum assured")
    ws.cell(row=r + 2, column=2, value=f"=B{r}*1000").number_format = "0.000"
    ws.cell(row=r + 3, column=1, value="Annuity of 100,000 a year for life, MXN")
    ws.cell(row=r + 3, column=2, value="=B15*100000").number_format = "#,##0"
    for rr in range(r, r + 4):
        ws.cell(row=rr, column=2).fill = OUTPUT_FILL
    ws.cell(
        row=r + 5,
        column=1,
        value=(
            "Net premiums only (mortality and interest, no expenses or profit): the gross premium "
            "from the profit test on the Rate_Card sheet adds acquisition cost, commission, "
            "maintenance and a target return."
        ),
    )
    ws.cell(row=r + 5, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + 5, start_column=1, end_row=r + 5, end_column=2)
    ws.row_dimensions[r + 5].height = 62


def _values_sheet(wb: Workbook, name: str, df: pd.DataFrame, note: str, formats=None) -> None:
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = note, Font(italic=True)
    columns = list(df.columns)
    _header(ws, columns, row=3)
    for i, record in enumerate(df.itertuples(index=False), start=4):
        for j, value in enumerate(record, start=1):
            if isinstance(value, np.generic):
                value = value.item()
            if isinstance(value, float) and not np.isfinite(value):
                value = None
            cell = ws.cell(row=i, column=j, value=value)
            fmt = (formats or {}).get(columns[j - 1])
            if fmt:
                cell.number_format = fmt


def _readme(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Mortality Table Builder: actuarial workbook", Font(bold=True, size=16)),
        (
            "Mexico's life tables with live commutation functions. Source: github.com/deadmigue55/Mortality-Table-Builder",
            None,
        ),
        ("", None),
        ("How to use", Font(bold=True, size=12)),
        (
            "1. Open the Calculator sheet. Blue cells are inputs: table, issue age, interest rate, term, sum assured.",
            None,
        ),
        (
            "2. Table_M, Table_F and Table_CNSF hold qx and the formulas for lx, dx, Dx, Nx, Cx and Mx. Click any cell to audit it.",
            None,
        ),
        (
            "3. Change the interest rate and every commutation column, premium and annuity moves with it.",
            None,
        ),
        (
            "4. Rate_Card, Excess_Mortality, Annuity_Stress and Table_Comparison hold results of the full pipeline as values.",
            None,
        ),
        ("", None),
        ("What the numbers are", Font(bold=True, size=12)),
        (
            "Tables: INEGI registered deaths and CONAPO population, completed for late registration and graduated (P-spline), year 2023.",
            None,
        ),
        (
            "CNSF 2013: the regulatory mortality table in force (mixed-sex). Net premiums here are mortality and interest only.",
            None,
        ),
        (
            "Rate_Card and Annuity_Stress use illustrative expense, commission and lapse assumptions and a synthetic book: methods, not a quote.",
            None,
        ),
        ("", None),
        ("Audit", Font(bold=True, size=12)),
        (
            "This workbook is recomputed by an independent formula engine in the project's tests and compared with the Python functions.",
            None,
        ),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True)
        if font:
            c.font = font


def build_workbook(con: sqlite3.Connection, defaults: dict | None = None) -> Workbook:
    """The workbook, from a connection to the warehouse or to the demo bundle."""
    wb = Workbook()
    _readme(wb)
    _calculator(wb, defaults)
    for (sheet, label), key in zip(TABLES, ("M", "F", "cnsf"), strict=True):
        _table_sheet(wb, sheet, label, _qx_curve(con, key))
    premiums = pd.read_sql(
        "SELECT sex AS Sex, issue_age AS [Issue age], term AS Term, net_premium_per_1000 AS "
        "[Net per 1000], gross_premium_per_1000 AS [Gross per 1000], gross_premium AS [Gross "
        "premium MXN], loading_over_net AS [Loading over net], irr AS IRR, year_1_profit AS "
        "[Year-1 profit], payback_year AS [Payback year] FROM pricing_premiums ORDER BY sex, term, "
        "issue_age",
        con,
    )
    _values_sheet(
        wb,
        "Rate_Card",
        premiums,
        "Gross premium of a level term policy, sum assured 1,000,000, solved for a 10% return on "
        "CNSF 2013 (illustrative expenses and lapses).",
        {
            "Net per 1000": "0.000",
            "Gross per 1000": "0.000",
            "Gross premium MXN": "#,##0",
            "Loading over net": "0%",
            "IRR": "0.0%",
            "Year-1 profit": "#,##0",
        },
    )
    excess = pd.read_sql(
        "SELECT year AS Year, observed AS Observed, expected_central AS Expected, excess AS "
        "[Excess deaths], excess_low AS Low, excess_high AS High, p_score AS [P-score], "
        "covid_coded AS [Certified COVID-19] FROM excess_summary WHERE sex = 'T' ORDER BY year",
        con,
    )
    _values_sheet(
        wb,
        "Excess_Mortality",
        excess,
        "Observed against expected deaths, both sexes; the central estimate is the median of four baselines.",
        {
            c: "#,##0"
            for c in ("Observed", "Expected", "Excess deaths", "Low", "High", "Certified COVID-19")
        }
        | {"P-score": "0.0%"},
    )
    stress = pd.read_sql(
        "SELECT label AS Basis, reserve_M AS [Reserve men], reserve_F AS [Reserve women], "
        "reserve_T AS [Reserve total], change_vs_central AS [Against central] FROM stress_longevity",
        con,
    )
    _values_sheet(
        wb,
        "Annuity_Stress",
        stress,
        "Reserve of an illustrative block of 20,000 pensioners (synthetic) on each mortality basis, annuity-due at 4%.",
        {
            "Reserve men": "#,##0",
            "Reserve women": "#,##0",
            "Reserve total": "#,##0",
            "Against central": "+0.0%;-0.0%",
        },
    )
    comparison = pd.read_sql(
        'SELECT source AS Source, "table" AS [Table], sex AS Sex, e0 AS e0, e65 AS e65, a_due_65 '
        "AS [Annuity at 65], A_40 AS [A(40)], P_term_20_at_40 AS [20-year term premium at 40] "
        'FROM life_table_comparison ORDER BY source, "table", sex',
        con,
    )
    _values_sheet(
        wb,
        "Table_Comparison",
        comparison,
        "Mexican population tables against regulatory and SOA tables, at 4%.",
        {
            "e0": "0.0",
            "e65": "0.0",
            "Annuity at 65": "0.00",
            "A(40)": "0.0000",
            "20-year term premium at 40": "0.00000",
        },
    )
    return wb


def workbook_bytes(con: sqlite3.Connection) -> bytes:
    buffer = io.BytesIO()
    build_workbook(con).save(buffer)
    return buffer.getvalue()


def write_workbook(con: sqlite3.Connection, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(con).save(path)
    return path
